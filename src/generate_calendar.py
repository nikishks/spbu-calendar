import json
import re
from datetime import datetime, date, timedelta
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
import openpyxl
from icalendar import Calendar, Event


ROOT = Path(__file__).resolve().parents[1]

CONFIG = json.loads(
    (ROOT / "config.json").read_text(encoding="utf-8")
)

API_EXCEL = "https://timetable.spbu.ru/StudentGroupEvents/ExcelWeek"


def norm(value):
    value = str(value or "")
    value = value.replace("\xa0", " ")
    value = value.strip().lower()
    value = re.sub(r"\s+", " ", value)
    return value


def is_selected_subject(subject):
    """
    Оставляем:
    - все обычные предметы;
    - только выбранные элективы.

    Убираем:
    - все остальные элективы;
    - все факультативы.
    """

    s = norm(subject)

    if not s:
        return False

    # Факультативы никогда не добавляем.
    if s.startswith("факультатив"):
        return False

    # Элективы.
    if s.startswith("электив"):
        selected_electives = [
            norm(x)
            for x in CONFIG.get("selected_electives", [])
        ]

        for elective in selected_electives:
            if elective in s:
                return True

        return False

    # Всё остальное — обычные предметы.
    return True


def get_sheet(group_id, monday):
    session = requests.Session()

    session.cookies.set(
        "_culture",
        "ru",
        domain="timetable.spbu.ru"
    )

    response = session.get(
        API_EXCEL,
        params={
            "studentGroupId": group_id,
            "weekMonday": monday.isoformat()
        },
        timeout=60,
    )

    response.raise_for_status()

    workbook = openpyxl.load_workbook(
        BytesIO(response.content),
        data_only=True
    )

    if not workbook.sheetnames:
        raise RuntimeError(
            "СПбГУ вернул Excel без листов."
        )

    return workbook[workbook.sheetnames[0]]


def parse_time(value):
    if not value:
        return None

    value = str(value)
    value = value.replace("—", "-")
    value = value.replace("–", "-")
    value = value.strip()

    parts = [
        x.strip()
        for x in value.split("-")
    ]

    if len(parts) != 2:
        return None

    try:
        start = datetime.strptime(
            parts[0],
            "%H:%M"
        ).time()

        end = datetime.strptime(
            parts[1],
            "%H:%M"
        ).time()

        return start, end

    except ValueError:
        return None


def parse_day(value, monday):
    """
    СПбГУ может отдавать:

        вторник
        8 сентября

    Поэтому дата определяется относительно недели.
    """

    if not value:
        return None

    text = str(value)
    text = text.replace("\n", " ")
    text = norm(text)

    weekdays = {
        "понедельник": 0,
        "вторник": 1,
        "среда": 2,
        "четверг": 3,
        "пятница": 4,
        "суббота": 5,
        "воскресенье": 6,
    }

    weekday = None

    for name, number in weekdays.items():
        if name in text:
            weekday = number
            break

    if weekday is None:
        return None

    match = re.search(
        r"\b(\d{1,2})\s+"
        r"(января|февраля|марта|апреля|мая|июня|"
        r"июля|августа|сентября|октября|ноября|декабря)\b",
        text
    )

    if not match:
        return None

    day = int(match.group(1))

    months = {
        "января": 1,
        "февраля": 2,
        "марта": 3,
        "апреля": 4,
        "мая": 5,
        "июня": 6,
        "июля": 7,
        "августа": 8,
        "сентября": 9,
        "октября": 10,
        "ноября": 11,
        "декабря": 12,
    }

    month = months[match.group(2)]

    year = monday.year

    # Если неделя пересекает Новый год.
    if month == 1 and monday.month == 12:
        year += 1

    try:
        result = date(
            year,
            month,
            day
        )
    except ValueError:
        return None

    # Проверяем, что дата соответствует дню недели.
    if result.weekday() != weekday:
        return None

    return result


def cell_is_struck(cell):
    """
    Возвращает True, если текст в ячейке Excel зачёркнут.
    """

    try:
        return bool(cell.font and cell.font.strike)
    except Exception:
        return False


def row_is_cancelled(row):
    """
    СПбГУ отмечает отменённые занятия зачёркиванием.

    Проверяем время, название, аудиторию и преподавателя.
    Первая колонка содержит день недели, поэтому её не учитываем.
    """

    important_cells = row[1:5]

    for cell in important_cells:
        if cell.value is None:
            continue

        if cell_is_struck(cell):
            return True

    return False


def parse_sheet(sheet, monday):
    result = []

    current_day = None

    for row in sheet.iter_rows(
        min_row=1,
        max_row=500,
        max_col=5
    ):
        values = []

        for cell in row:
            value = cell.value

            if value is not None:
                value = str(value)
                value = value.replace("\xa0", " ")
                value = value.strip()

            values.append(value)

        if not any(values):
            continue

        first = values[0] or ""

        # Если в первой ячейке указан день,
        # запоминаем его для следующих строк.
        parsed_day = parse_day(
            first,
            monday
        )

        if parsed_day:
            current_day = parsed_day

        # Структура:
        # 0 — день
        # 1 — время
        # 2 — название
        # 3 — место
        # 4 — преподаватель

        if not current_day:
            continue

        if len(values) < 5:
            continue

        time_value = values[1]
        subject = values[2]

        if not time_value or not subject:
            continue

        parsed_time = parse_time(
            time_value
        )

        if not parsed_time:
            continue

        start, end = parsed_time

        # ----------------------------------------
        # ПРОВЕРКА НА ОТМЕНУ
        # ----------------------------------------

        if row_is_cancelled(row):
            print(
                "Отменённое занятие пропущено: "
                f"{current_day} "
                f"{start.strftime('%H:%M')} "
                f"{subject}"
            )
            continue

        # ----------------------------------------
        # ФИЛЬТРАЦИЯ ПРЕДМЕТОВ
        # ----------------------------------------

        if not is_selected_subject(subject):
            continue

        place = values[3] or ""
        lecturer = values[4] or ""

        result.append(
            (
                current_day,
                start,
                end,
                subject,
                place,
                lecturer
            )
        )

    return result


def monday_on_or_before(d):
    return d - timedelta(
        days=d.weekday()
    )


def load_schedule():
    group_id = int(
        CONFIG["student_group_id"]
    )

    start = date.fromisoformat(
        CONFIG["start_date"]
    )

    end = start + timedelta(
        days=30 * int(
            CONFIG.get(
                "months_ahead",
                8
            )
        )
    )

    monday = monday_on_or_before(
        start
    )

    all_events = []

    while monday <= end:
        print(
            f"Загрузка недели: {monday}"
        )

        sheet = get_sheet(
            group_id,
            monday
        )

        week_events = parse_sheet(
            sheet,
            monday
        )

        print(
            f"Найдено подходящих занятий: "
            f"{len(week_events)}"
        )

        all_events.extend(
            week_events
        )

        monday += timedelta(
            days=7
        )

    # Убираем полные дубликаты.
    unique = {}

    for event in all_events:
        unique[event] = event

    return sorted(
        unique.values(),
        key=lambda x: (
            x[0],
            x[1],
            x[3]
        )
    )


def make_uid(
    event_date,
    start,
    subject
):
    """
    UID должен оставаться стабильным.

    Поэтому НЕ включаем:
    - аудиторию;
    - преподавателя;
    - время окончания.

    Если поменяется аудитория или преподаватель,
    календарь должен обновить существующее событие,
    а не создать второе.
    """

    uid_base = (
        f"{event_date.isoformat()}-"
        f"{start.strftime('%H-%M')}-"
        f"{norm(subject)}"
    )

    uid_clean = re.sub(
        r"[^0-9A-Za-zА-Яа-яЁё_.-]+",
        "-",
        uid_base
    ).strip("-")

    return (
        uid_clean
        + "@spbu-calendar"
    )


def make_ics(events):
    timezone_name = CONFIG.get(
        "timezone",
        "Europe/Moscow"
    )

    timezone = ZoneInfo(
        timezone_name
    )

    calendar = Calendar()

    calendar.add(
        "prodid",
        "-//SPbU Apple Calendar//RU//"
    )

    calendar.add(
        "version",
        "2.0"
    )

    calendar.add(
        "CALSCALE",
        "GREGORIAN"
    )

    calendar.add(
        "METHOD",
        "PUBLISH"
    )

    calendar.add(
        "X-WR-CALNAME",
        CONFIG.get(
            "calendar_name",
            "СПбГУ"
        )
    )

    calendar.add(
        "X-WR-CALDESC",
        CONFIG.get(
            "calendar_description",
            "Расписание СПбГУ"
        )
    )

    calendar.add(
        "X-WR-TIMEZONE",
        timezone_name
    )

    now = datetime.now(timezone)

    for (
        event_date,
        start,
        end,
        subject,
        place,
        lecturer
    ) in events:

        event = Event()

        start_dt = datetime.combine(
            event_date,
            start,
            tzinfo=timezone
        )

        end_dt = datetime.combine(
            event_date,
            end,
            tzinfo=timezone
        )

        uid = make_uid(
            event_date,
            start,
            subject
        )

        event.add(
            "uid",
            uid
        )

        event.add(
            "dtstamp",
            now
        )

        event.add(
            "last-modified",
            now
        )

        event.add(
            "sequence",
            0
        )

        event.add(
            "dtstart",
            start_dt
        )

        event.add(
            "dtend",
            end_dt
        )

        event.add(
            "summary",
            subject
        )

        description = []

        if lecturer:
            description.append(
                f"Преподаватель: {lecturer}"
            )

        if place:
            description.append(
                f"Аудитория: {place}"
            )

        description.append(
            "Источник: timetable.spbu.ru"
        )

        event.add(
            "description",
            "\n".join(description)
        )

        if place:
            event.add(
                "location",
                place
            )

        calendar.add_component(
            event
        )

    return calendar.to_ical()


def main():
    events = load_schedule()

    output = (
        ROOT
        / "docs"
        / "schedule.ics"
    )

    output.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    output.write_bytes(
        make_ics(events)
    )

    print()
    print(
        f"Создано событий: {len(events)}"
    )

    print(
        f"Файл: {output}"
    )


if __name__ == "__main__":
    main()
